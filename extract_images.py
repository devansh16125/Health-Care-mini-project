import pptx
import os

prs = pptx.Presentation(r"E:\HealthCare_Presentation.pptx")
img_dir = r"C:\Users\devan\.gemini\antigravity-ide\scratch\alzheimer-explainable-ml\presentation\extracted_user_images"
os.makedirs(img_dir, exist_ok=True)

img_count = 0
for s_idx, slide in enumerate(prs.slides, 1):
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            ext = img.ext
            filename = f"slide_{s_idx:02d}_img_{sh_idx}_{shape.name}.{ext}"
            filepath = os.path.join(img_dir, filename)
            with open(filepath, "wb") as f:
                f.write(img.blob)
            img_count += 1
            print(f"Extracted: {filename} ({len(img.blob)/1024:.1f} KB)")

print(f"Total extracted: {img_count} images to {img_dir}")
