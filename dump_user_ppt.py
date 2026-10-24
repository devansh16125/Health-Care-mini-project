import pptx
import sys

prs = pptx.Presentation(r"E:\HealthCare_Presentation.pptx")

with open(r"user_ppt_content.txt", "w", encoding="utf-8") as out:
    out.write(f"Total slides in E:\\HealthCare_Presentation.pptx: {len(prs.slides)}\n")

    for idx, slide in enumerate(prs.slides, 1):
        out.write("=" * 80 + "\n")
        out.write(f"SLIDE {idx}\n")
        out.write("-" * 80 + "\n")
        
        # Speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.write(f"[SPEAKER NOTES]:\n{notes}\n\n")
                
        # Shapes
        for s_idx, shape in enumerate(slide.shapes):
            name = shape.name
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    out.write(f"  [Text Shape {s_idx} ({name})]:\n{text}\n\n")
            elif shape.has_table:
                table = shape.table
                out.write(f"  [Table {s_idx}]: {len(table.rows)} rows x {len(table.columns)} cols\n")
                for r in table.rows:
                    row_txt = [c.text.strip().replace("\n", " ") for c in r.cells]
                    out.write(f"    | {' | '.join(row_txt)} |\n")
                out.write("\n")
            elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                out.write(f"  [Picture {s_idx} ({name})]: image bytes: {len(shape.image.blob)} bytes, format: {shape.image.content_type}\n")

print("Dumped successfully to user_ppt_content.txt")
